"""Extract the Facebook Messenger game chats ("Oral Histories") into the private canon repo.

Usage: python scripts/oral_histories.py [--notes-dir DIR]

Every subfolder of the notes directory that holds a message_1.json is one chat; its output folder
is the chat title, slugified.

Re-runnable. The Messenger export folders are only read. Everything this writes goes under
canon/oral-histories/, which is private and never pushed: the transcripts hold out-of-game
conversation between real people. Articles cite them as "Oral Histories: <chat>, YYYY-MM-DD".

Writes, per chat:
  <chat>/YYYY-MM.md          the transcript, one line per message, local time
  <chat>/attachments/<id>.md text of every docx/pdf/md/txt file posted, with who, when and context
  <chat>/media-index.csv     every photo, gif, video and file, with size and neighbouring messages
and STATS.md across all chats.
"""
import csv
import datetime
import json
import os
import re
import sys
from collections import Counter, defaultdict

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'canon', 'oral-histories')
NOTES_DIR = os.path.join(os.path.expanduser('~'), 'Desktop', 'D&D', 'Notes')

MEDIA_KEYS = ('photos', 'gifs', 'videos', 'files')
TEXT_EXT = {'.docx', '.pdf', '.md', '.txt'}
REACTED_RE = re.compile(r'^.{1,60} reacted .{1,8} to your message\s*$')
CONTEXT = 3


def fix(s):
    """Messenger exports store UTF-8 bytes as latin-1 code points."""
    if not s:
        return ''
    try:
        return s.encode('latin-1').decode('utf-8')
    except (UnicodeEncodeError, UnicodeDecodeError):
        return s


def local(ms):
    return datetime.datetime.fromtimestamp(ms / 1000)


def write(path, text):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, 'w', encoding='utf-8', newline='\n') as f:
        f.write(text)


def load_chat(folder):
    messages, title, participants = [], None, set()
    for name in sorted(os.listdir(folder)):
        if re.match(r'message_\d+\.json$', name):
            with open(os.path.join(folder, name), encoding='utf-8') as f:
                j = json.load(f)
            title = fix(j.get('title'))
            participants.update(fix(p['name']) for p in j.get('participants', []))
            messages.extend(j['messages'])
    messages = [m for m in messages if not m.get('is_unsent')]
    messages.sort(key=lambda m: m['timestamp_ms'])
    return title, sorted(participants), messages


def media_of(m, folder):
    """(kind, local path, basename) for every attachment on a message."""
    out = []
    for key in MEDIA_KEYS:
        for item in m.get(key, []) or []:
            base = os.path.basename(item.get('uri', ''))
            out.append((key, os.path.join(folder, key, base), base))
    return out


def text_of(m):
    c = fix(m.get('content'))
    if c and REACTED_RE.match(c):
        return ''
    return c


def line_for(m, folder):
    t = local(m['timestamp_ms']).strftime('%Y-%m-%d %H:%M')
    parts = []
    c = text_of(m)
    if c:
        parts.append(c.replace('\r', '').replace('\n', ' / '))
    for kind, _, base in media_of(m, folder):
        parts.append('{%s: %s/%s}' % (kind.rstrip('s'), kind, base))
    share = m.get('share') or {}
    if share.get('link') and share['link'] not in (c or ''):
        parts.append('{link: %s}' % share['link'])
    if not parts:
        return None
    return '[%s] %s: %s' % (t, fix(m.get('sender_name')), ' '.join(parts))


def neighbours(messages, i):
    before, after = [], []
    j = i - 1
    while j >= 0 and len(before) < CONTEXT:
        c = text_of(messages[j])
        if c:
            before.insert(0, '%s: %s' % (fix(messages[j].get('sender_name')), c[:300]))
        j -= 1
    j = i + 1
    while j < len(messages) and len(after) < CONTEXT:
        c = text_of(messages[j])
        if c:
            after.append('%s: %s' % (fix(messages[j].get('sender_name')), c[:300]))
        j += 1
    return before, after


def extract_text(path):
    """Return (text, note). Never raises."""
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext in ('.md', '.txt'):
            with open(path, encoding='utf-8', errors='replace') as f:
                return f.read(), ''
        if ext == '.docx':
            import docx
            d = docx.Document(path)
            lines = [p.text for p in d.paragraphs]
            for table in d.tables:
                for row in table.rows:
                    lines.append(' | '.join(cell.text.strip() for cell in row.cells))
            return '\n'.join(lines), ''
        if ext == '.pdf':
            import fitz
            doc = fitz.open(path)
            pages = [page.get_text() for page in doc]
            text = '\n\n'.join('--- page %d ---\n%s' % (n + 1, p) for n, p in enumerate(pages))
            chars = sum(len(p.strip()) for p in pages)
            note = '' if chars > 50 * len(pages) else 'little or no text layer (%d pages, %d chars): likely scanned or image-only' % (len(pages), chars)
            return text, note
    except Exception as e:  # noqa: BLE001 - log every failure, keep going
        return '', 'extraction failed: %s' % e
    return '', 'unsupported type'


def image_size(path):
    try:
        from PIL import Image
        with Image.open(path) as im:
            return '%dx%d' % im.size
    except Exception:  # noqa: BLE001
        return ''


def slugify(title):
    title = re.sub(r',?\s+LLC$', '', title)
    return re.sub(r'[^a-z0-9]+', '-', title.lower()).strip('-')


def process(folder, stats):
    title, participants, messages = load_chat(folder)
    slug = slugify(title)
    out = os.path.join(OUT, slug)

    by_month = defaultdict(list)
    for m in messages:
        line = line_for(m, folder)
        if line:
            by_month[local(m['timestamp_ms']).strftime('%Y-%m')].append(line)
    for month, lines in by_month.items():
        head = '# %s - %s\n\nOral Histories: %s. Private transcript, never published.\n\n' % (title, month, title)
        write(os.path.join(out, month + '.md'), head + '\n'.join(lines) + '\n')

    rows, attachments, missing = [], [], []
    for i, m in enumerate(messages):
        for kind, path, base in media_of(m, folder):
            exists = os.path.exists(path)
            if not exists:
                missing.append('%s/%s' % (kind, base))
            before, after = neighbours(messages, i)
            when = local(m['timestamp_ms']).strftime('%Y-%m-%d %H:%M')
            sender = fix(m.get('sender_name'))
            caption = text_of(m)
            rows.append({
                'file': '%s/%s' % (kind, base), 'kind': kind.rstrip('s'),
                'ext': os.path.splitext(base)[1].lower(),
                'bytes': os.path.getsize(path) if exists else '',
                'pixels': image_size(path) if exists and kind in ('photos', 'gifs') else '',
                'date': when, 'sender': sender, 'caption': caption,
                'before': ' || '.join(before), 'after': ' || '.join(after),
            })
            if kind == 'files' and os.path.splitext(base)[1].lower() in TEXT_EXT and exists:
                text, note = extract_text(path)
                stem = os.path.splitext(base)[0]
                first = next((ln.strip() for ln in text.splitlines()
                              if ln.strip() and not ln.startswith('--- page')), '')
                attachments.append((when, sender, '%s/%s' % (kind, base), first[:120], note, len(text)))
                body = ['# Attachment %s' % base, '',
                        '- Chat: %s (Oral Histories)' % title,
                        '- Posted: %s by %s' % (when, sender),
                        '- Source: `%s`' % path]
                if caption:
                    body.append('- Caption: %s' % caption)
                if note:
                    body.append('- Note: %s' % note)
                body += ['', '## Context before', ''] + ['> ' + b for b in before]
                body += ['', '## Context after', ''] + ['> ' + a for a in after]
                body += ['', '## Text', '', text]
                write(os.path.join(out, 'attachments', stem + '.md'), '\n'.join(body) + '\n')

    os.makedirs(out, exist_ok=True)
    with open(os.path.join(out, 'media-index.csv'), 'w', encoding='utf-8', newline='') as f:
        w = csv.DictWriter(f, fieldnames=['file', 'kind', 'ext', 'bytes', 'pixels', 'date', 'sender',
                                          'caption', 'before', 'after'])
        w.writeheader()
        w.writerows(rows)

    stats.append({'slug': slug, 'title': title, 'participants': participants, 'messages': messages,
                  'months': by_month, 'media': rows, 'attachments': attachments, 'missing': missing,
                  'folder': folder})


def render_stats(stats):
    out = ['# Oral Histories - extraction stats', '',
           'Generated by `scripts/oral_histories.py`. Private: never published.', '']
    out += ['| Chat | Folder | Messages | Text lines | Span | Media | Documents | Missing files |',
            '|---|---|---|---|---|---|---|---|']
    for s in stats:
        ms = s['messages']
        span = '%s to %s' % (local(ms[0]['timestamp_ms']).date(), local(ms[-1]['timestamp_ms']).date()) if ms else ''
        lines = sum(len(v) for v in s['months'].values())
        out.append('| %s | `%s/` | %d | %d | %s | %d | %d | %d |' % (
            s['title'], s['slug'], len(ms), lines, span, len(s['media']), len(s['attachments']), len(s['missing'])))
    for s in stats:
        out += ['', '## %s' % s['title'], '', 'Participants: ' + ', '.join(s['participants']), '']
        kinds = Counter((r['kind'], r['ext']) for r in s['media'])
        out.append('Media: ' + ', '.join('%s %s: %d' % (k, e, n) for (k, e), n in sorted(kinds.items())))
        if s['missing']:
            out.append('')
            out.append('Missing from the export folder: ' + ', '.join(s['missing'][:20]) +
                       (' ...' if len(s['missing']) > 20 else ''))
        if s['attachments']:
            out += ['', '### Documents', '', '| Posted | By | File | Chars | First line | Note |',
                    '|---|---|---|---|---|---|']
            for when, sender, f, first, note, n in s['attachments']:
                out.append('| %s | %s | `%s` | %d | %s | %s |' % (
                    when, sender, f, n, first.replace('|', '/'), note))
        per_year = defaultdict(Counter)
        for m in s['messages']:
            per_year[local(m['timestamp_ms']).year][fix(m.get('sender_name'))] += 1
        people = [p for p, _ in Counter(fix(m.get('sender_name')) for m in s['messages']).most_common()]
        out += ['', '### Messages per year per person', '',
                '| Year | ' + ' | '.join(people) + ' |', '|---' * (len(people) + 1) + '|']
        for y in sorted(per_year):
            out.append('| %d | ' % y + ' | '.join(str(per_year[y][p]) for p in people) + ' |')
        days = Counter(local(m['timestamp_ms']).strftime('%Y-%m-%d (%a)') for m in s['messages'])
        out += ['', '### Busiest days (likely session nights)', '']
        out += ['- %s: %d messages' % (d, n) for d, n in days.most_common(25)]
    return '\n'.join(out) + '\n'


def main():
    notes = NOTES_DIR
    if '--notes-dir' in sys.argv:
        notes = sys.argv[sys.argv.index('--notes-dir') + 1]
    if not os.path.isdir(os.path.join(ROOT, 'canon')):
        sys.exit('canon/ is missing: this script only writes into the private canon repo')
    stats = []
    for name in sorted(os.listdir(notes)):
        folder = os.path.join(notes, name)
        if os.path.isdir(folder) and os.path.exists(os.path.join(folder, 'message_1.json')):
            process(folder, stats)
            print('%s: %d messages' % (stats[-1]['slug'], len(stats[-1]['messages'])))
    write(os.path.join(OUT, 'STATS.md'), render_stats(stats))
    print('wrote', OUT)


if __name__ == '__main__':
    main()
